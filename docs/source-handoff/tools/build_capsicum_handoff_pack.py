from __future__ import annotations

import csv
import hashlib
import json
import os
import shutil
import textwrap
import zipfile
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

ROOT = Path('/mnt/data/capsicum-breeding-platform-agent-handoff/capsicum-breeding-intelligence-platform')
SOURCE_ROOT = Path('/mnt/data')


def write(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    normalized = textwrap.dedent(content).strip() + '\n'
    path.write_text(normalized, encoding='utf-8')


def dump_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, indent=2, ensure_ascii=False) + '\n', encoding='utf-8')


# Clean generated subdirectories while preserving scaffold root.
for name in ['docs', 'source-pack', 'data', 'contracts', 'starter-repo-blueprint', 'checklists', 'tools']:
    p = ROOT / name
    if p.exists():
        shutil.rmtree(p)
    p.mkdir(parents=True, exist_ok=True)

# Source assets
source_assets = [
    'CAPSICUM_BREEDING_INTELLIGENCE_PLATFORM_IMPLEMENTATION_PLAN.md',
    'CAPSICUM_GENETICS_DATASET_README_v0_1.md',
    'capsicum_genetics_evidence_catalog_v0_1.xlsx',
    'capsicum_locus_catalog_v0_1.csv',
    'capsicum_sources_v0_1.csv',
    'build_capsicum_catalog.py',
]
for name in source_assets:
    src = SOURCE_ROOT / name
    if not src.exists():
        raise FileNotFoundError(src)
    shutil.copy2(src, ROOT / 'source-pack' / name)

# Export every workbook sheet to CSV and JSON for deterministic agent ingestion.
wb_path = SOURCE_ROOT / 'capsicum_genetics_evidence_catalog_v0_1.xlsx'
wb = load_workbook(wb_path, read_only=True, data_only=False)
workbook_manifest: dict[str, Any] = {'source_file': wb_path.name, 'sheets': []}
for ws in wb.worksheets:
    rows = [[cell.value for cell in row] for row in ws.iter_rows()]
    safe = ''.join(ch if ch.isalnum() or ch in '-_' else '_' for ch in ws.title)
    csv_path = ROOT / 'data' / f'{safe}.csv'
    with csv_path.open('w', newline='', encoding='utf-8') as f:
        writer = csv.writer(f)
        writer.writerows(rows)
    json_path = ROOT / 'data' / f'{safe}.json'
    dump_json(json_path, rows)
    workbook_manifest['sheets'].append({
        'name': ws.title,
        'rows': ws.max_row,
        'columns': ws.max_column,
        'csv': f'data/{safe}.csv',
        'json': f'data/{safe}.json',
    })
dump_json(ROOT / 'data' / 'workbook-export-manifest.json', workbook_manifest)

write(ROOT / 'START_HERE.md', r'''
# Start Here — Capsicum Breeding Intelligence Platform

This archive is a self-contained implementation handoff for a **web-based Capsicum breeding intelligence platform**. It is designed to be given to a capable autonomous coding agent or an engineering team.

## Required reading order

1. `SOURCE_PACK_READINESS.md`
2. `EXECUTION_PROMPT.md`
3. `source-pack/CAPSICUM_BREEDING_INTELLIGENCE_PLATFORM_IMPLEMENTATION_PLAN.md`
4. `docs/00_PRODUCT_CHARTER.md`
5. `docs/01_SYSTEM_ARCHITECTURE.md`
6. `docs/02_WEB_APPLICATION_UX_SPEC.md`
7. `docs/03_DATA_MODEL_AND_MIGRATIONS.md`
8. `docs/04_SIMULATION_ENGINE_SPEC.md`
9. `docs/05_AI_AND_RAG_SPEC.md`
10. `docs/06_PHENOTYPE_VISION_SPEC.md`
11. `WORKSTREAM_CARDS.md`
12. `VALIDATION_LEDGER.md`

## How to use this pack

Give the entire extracted directory to the coding agent. Use the contents of `EXECUTION_PROMPT.md` as the agent's primary instruction. The agent must treat all other files as binding source material.

The agent is expected to create a new repository from this specification. The archive is not the completed application and does not pretend that unsupported quantitative genetics models can be made scientifically accurate without training data.

## Authoritative hierarchy

When documents conflict, use this order:

1. Scientific safety, prohibited claims, evidence eligibility, and stop conditions.
2. `EXECUTION_PROMPT.md`.
3. The full implementation plan in `source-pack/`.
4. Domain specifications under `docs/`.
5. Machine-readable schemas under `contracts/`.
6. Starter repository blueprint.
7. Informational notes and examples.

Do not resolve conflicts by weakening scientific integrity, access controls, auditability, test requirements, or uncertainty disclosure.

## Deliverable expected from the coding agent

A complete, runnable, documented monorepo containing:

- responsive Next.js web application;
- PostgreSQL database and migrations;
- exact and advanced inheritance simulation modules;
- scientific evidence catalog and review workflows;
- germplasm, seed-lot, plant, cross, pedigree, experiment, phenotype, and selection workflows;
- bounded AI research and breeding copilot;
- phenotype image capture, deterministic measurement, annotation, and model-evaluation infrastructure;
- Python and R worker boundaries for scientific workloads;
- local Docker Compose deployment;
- automated tests, CI, observability, backups, and restore documentation;
- seeded demonstration workspace using the supplied Capsicum catalog.

## Release truth

The first application can be functionally complete while some scientific capabilities remain explicitly **research-only** or **blocked pending data**. A missing validated model must result in abstention, not fabricated accuracy.
''')

write(ROOT / 'README.md', r'''
# Capsicum Breeding Intelligence Platform — Agent Handoff Pack

## Purpose

This package provides the scientific evidence seed, architecture, product requirements, web UX specification, implementation sequence, machine-readable contracts, validation gates, and autonomous-agent execution prompt required to create a serious Capsicum breeding application.

## Product promise

The platform will help breeders:

- record germplasm, seed lots, plants, controlled crosses, harvests, progeny, and pedigrees;
- calculate defensible inheritance probabilities for approved genetic models;
- simulate linkage, maternal inheritance, epistasis, uncertain genotypes, and selection programs;
- calculate practical target-recovery population sizes;
- collect standardized phenotype and environmental observations;
- capture and analyze plant and fruit imagery through controlled protocols;
- discover, review, and curate scholarly evidence with AI assistance;
- explain results without hiding assumptions, uncertainty, applicability, or source provenance;
- build validated quantitative, genomic, and vision models only when adequate data exists.

## Explicit nonclaims

The platform must not imply that it can universally or exactly predict:

- Scoville heat units from parent names or values;
- flavor;
- yield;
- fruit size or wall thickness without a validated model;
- disease resistance without pathogen context;
- environmental adaptation through arbitrary multipliers;
- a cultivar's genotype without direct evidence;
- gene-editing outcomes without a validated molecular model.

## Contents

- `EXECUTION_PROMPT.md`: one-shot master coding-agent prompt.
- `WORKSTREAM_CARDS.md`: bounded multi-agent workstreams and merge order.
- `source-pack/`: original evidence workbook, CSVs, source registry, and full implementation plan.
- `docs/`: implementation specifications.
- `contracts/`: JSON schemas and canonical contracts.
- `data/`: workbook sheets exported to CSV and JSON.
- `starter-repo-blueprint/`: expected repository topology and environment baseline.
- `checklists/`: release, scientific, security, and UI completion gates.
- `FINAL_REPORT_TEMPLATE.md`: mandatory agent completion report.

## Readiness

**Conditionally implementation-ready.** The application architecture, exact genetics engine, catalog governance, breeding ledger, AI boundaries, visual-analysis workflow, and web interface can be implemented. Production quantitative/genomic predictions and learned vision claims remain gated by representative data and independent validation.
''')

write(ROOT / 'SOURCE_PACK_READINESS.md', r'''
# Source Pack Readiness

## Classification

- Project: Capsicum Breeding Intelligence Platform
- Execution mode: Greenfield product and scientific data platform
- Repository shape: Multi-app, multi-language monorepo
- Multi-agent execution: Supported and recommended
- Readiness state: **Conditionally ready**

## Evidence assessment

| Area | Present | Status |
|---|---:|---|
| Product outcome and users | Yes | Sufficient |
| Scientific evidence seed | Yes | 22 locus records, 22 claims, 30 sources; requires deeper allele-level curation |
| Scope and explicit non-goals | Yes | Binding |
| Architecture and package boundaries | Yes | Binding baseline |
| Web interface requirements | Yes | Binding |
| Data and provenance requirements | Yes | Binding |
| Simulation engine taxonomy | Yes | Exact engines ready; data-trained engines gated |
| AI role and prohibition boundaries | Yes | Binding |
| Vision workflow | Yes | Infrastructure ready; model performance gated by dataset |
| Security, permissions, and audit | Yes | Binding |
| Validation and release gates | Yes | Binding |
| Deployment and rollback | Yes | Local-first baseline defined |

## Conditions that do not block repository creation

The following must be implemented as configurable, feature-gated, or abstaining capabilities rather than treated as reasons to abandon the build:

- absent production AI provider credentials;
- absent S3 credentials in local development;
- absent learned phenotype models;
- absent genomic datasets;
- insufficient data for quantitative trait prediction;
- scientific claims awaiting review.

Use local emulators, mock providers, deterministic fixtures, feature flags, and explicit `not_validated` states.

## Conditions that block a production scientific claim

- no supporting evidence assertion;
- no approved rule version;
- parent genotype unknown but treated as verified;
- unsupported species or population applicability;
- missing pathogen isolate context for resistance claims;
- missing capture protocol or failed image-quality gate;
- no independent validation for learned or statistical models;
- generated results cannot be reproduced from stored inputs and versions.

## Required agent behavior

The coding agent may make implementation-level decisions that are consistent with this pack, but it may not weaken the scientific claim policy or silently fill data gaps. Any material conflict must be surfaced in the final report.
''')

write(ROOT / 'docs/00_PRODUCT_CHARTER.md', r'''
# Product Charter

## Problem

Pepper breeding records are commonly fragmented across notebooks, labels, spreadsheets, photographs, seed packets, and memory. Existing Punnett-square tools generally treat genetics as simple dominant/recessive pairs, omit provenance and uncertainty, and cannot support controlled breeding programs, evidence curation, visual phenotyping, or validated quantitative models.

## Primary users

- hobby and community pepper breeders;
- small seed companies;
- plant geneticists and breeding researchers;
- germplasm curators;
- greenhouse and field technicians;
- scientific reviewers and catalog administrators.

## Core jobs

1. Establish trustworthy biological identity and provenance.
2. Plan and record controlled crosses.
3. Simulate supported inheritance models.
4. Define target phenotypes and genotypes.
5. Determine grow-out population requirements.
6. Record standardized observations and images.
7. Compare predictions with observed progeny.
8. Curate and approve scientific evidence.
9. Train and validate models on accumulated data.
10. Export the complete breeding record in open formats.

## Product principles

- Evidence before prediction.
- Abstention before unsupported certainty.
- Exact mathematics where the biology permits it.
- Explicit uncertainty where inputs or models are uncertain.
- Human approval for scientific rule activation and genotype verification.
- Complete lineage and provenance.
- Web workflows usable by nontechnical operators.
- Local-first deployment with cloud portability.
- Data ownership and full export.

## Public-quality first release

The first release must support account/workspace management, evidence catalog import and review, germplasm and seed-lot management, plants, crosses, progeny, pedigrees, exact inheritance, unknown genotype distributions, target recovery calculations, observation capture, evidence-linked result pages, exports, audit history, backups, and a seeded demo workspace.

AI and vision infrastructure must be functional but must identify whether an output is deterministic, model-derived, AI-assisted, draft, reviewed, or unsupported.
''')

write(ROOT / 'docs/01_SYSTEM_ARCHITECTURE.md', r'''
# System Architecture

## Architectural style

Use a modular monorepo with a Next.js web control plane, shared TypeScript domain packages, PostgreSQL, S3-compatible object storage, and isolated Python/R scientific workers. Do not create a distributed microservice estate.

## Runtime topology

```text
Browser
  -> Next.js web application
       -> application/domain services
       -> exact TypeScript simulation engine
       -> PostgreSQL
       -> object storage
       -> PostgreSQL-backed job queue
            -> Python vision/statistics worker
            -> R breeding/statistics worker
       -> provider-agnostic AI gateway
```

## Recommended repository topology

```text
apps/
  web/                         Next.js App Router application
  worker-python/               FastAPI or worker process for vision/statistics
  worker-r/                    Plumber/CLI worker for AlphaSimR and R models
packages/
  auth/                        identity, workspace, roles, policies
  database/                    schema, migrations, repositories, transactions
  contracts/                   schemas, event/job contracts, generated types
  genetics-core/               pure exact inheritance engine
  genetics-advanced/           linkage, cytoplasmic, epistasis, uncertainty
  scientific-catalog/          evidence, rules, releases, eligibility
  breeding-domain/             germplasm, material identity, crosses, pedigree
  observation-domain/          studies, variables, observations, environment
  simulation-domain/           runs, snapshots, artifacts, model registry
  ai/                          bounded agents, retrieval, tools, evals
  vision/                      capture, image metadata, annotations, measurements
  jobs/                        queue contracts, retries, idempotency
  storage/                     object store abstraction and integrity
  observability/               logs, metrics, traces, audit events
  ui/                          design system and accessible components
  config/                      typed runtime configuration
  test-utils/                  fixtures, factories, reference crosses
infra/
  docker/
  compose/
  backup/
  observability/
docs/
```

## Dependency direction

- Domain packages must not import web UI.
- `genetics-core` must remain pure and deterministic with no database, network, AI, or React dependency.
- Catalog rules may call genetics primitives through compiled rule definitions; the engine must not query papers or LLMs.
- AI may call catalog and simulation APIs; catalog and simulation packages must not depend on AI.
- Workers consume versioned job contracts and return immutable artifacts.
- The web app owns authorization and orchestration; workers must not make access-control decisions from user-supplied fields.

## Technology baseline

At the date of this pack:

- Node.js 24 LTS for production.
- Next.js 16.2, using the latest compatible patch.
- AI SDK 6, using schema-validated tools and bounded loops.
- PostgreSQL 18, using the current supported minor release.
- MIAPPE 1.2 with compatibility mappings for 1.1.
- BrAPI 2.1 only for selected interoperability endpoints.

The agent must confirm compatible current patches before locking dependencies and record the final versions.

## Database approach

Use PostgreSQL-native constraints, transactions, JSONB only for genuinely flexible payloads, and SQL migrations as the source of truth. An ORM/query builder may be used, but it must not prevent partial indexes, exclusion constraints, recursive pedigree checks, vector search, or explicit locking.

## Job model

Use a PostgreSQL-backed queue initially. Every job must include:

- immutable input payload or content hash;
- workspace and actor context derived by the application;
- job type and schema version;
- worker and model version;
- idempotency key;
- retry policy;
- status and structured error;
- logs and artifact references;
- cancellation and timeout semantics.

## Deployment

Local and single-server operation uses Docker Compose. The design remains compatible with managed PostgreSQL and S3 storage, but Kubernetes is explicitly out of scope until load evidence justifies it.
''')

write(ROOT / 'docs/02_WEB_APPLICATION_UX_SPEC.md', r'''
# Web Application UX Specification

## Interface mandate

The product is a responsive web application. It must work well on desktop, tablet, and modern mobile browsers. The primary experience is not a developer console. It must guide growers and breeders who may have little genetics or software experience.

## Design standard

- WCAG 2.2 AA target.
- Keyboard-complete workflows.
- Semantic HTML and visible focus.
- Responsive layouts without horizontal page scrolling.
- Plain-language explanations with scientific details available progressively.
- Every destructive or scientifically authoritative action requires explicit confirmation.
- Empty, loading, partial, error, offline/degraded, and permission-denied states are designed—not omitted.
- Never communicate scientific state by color alone.

## Primary navigation

```text
Dashboard
Breeding
  Germplasm
  Seed Lots
  Plants
  Crosses
  Families and Progeny
  Pedigrees
Simulations
  New Simulation
  Saved Runs
  Targets and Selection Plans
Phenotyping
  Experiments
  Observation Sessions
  Images and Measurements
  Annotation Queue
Science
  Evidence Catalog
  Sources
  Candidate Claims
  Review Queue
AI Copilot
Reports and Exports
Administration
```

## Required route map

```text
/
/sign-in
/onboarding
/dashboard
/germplasm
/germplasm/new
/germplasm/[id]
/seed-lots
/seed-lots/[id]
/plants
/plants/[id]
/crosses
/crosses/new
/crosses/[id]
/families/[id]
/pedigrees/[materialId]
/simulations
/simulations/new
/simulations/[id]
/selection-plans
/selection-plans/[id]
/experiments
/experiments/[id]
/observations/session/[id]
/phenotypes/images
/phenotypes/images/[id]
/annotations
/catalog
/catalog/loci/[id]
/catalog/claims/[id]
/catalog/releases/[id]
/research/sources/[id]
/research/review
/ai
/reports
/settings/workspace
/settings/users
/admin/jobs
/admin/models
/admin/audit
/admin/system
```

## Dashboard

Show actionable operational information:

- planned crosses due today;
- pollinated flowers awaiting fruit-set review;
- seed lots approaching low inventory;
- plants needing observations;
- failed image-quality checks;
- simulation runs requiring assumption review;
- evidence claims awaiting review;
- jobs or model runs needing attention;
- recent selections and harvests.

Do not make vanity charts the primary dashboard content.

## Cross-planning workflow

1. Select a maternal parent and paternal parent.
2. Display species, seed lot, pedigree, known genotypes, assumed genotypes, and identity warnings.
3. Choose approved loci, traits, and target outcomes.
4. Review model applicability and excluded predictions.
5. Resolve or acknowledge unknown genotype states.
6. Run exact simulation or select an approved advanced model.
7. Present genotype, phenotype, and uncertainty distributions separately.
8. Calculate target recovery and adjusted planting requirements.
9. Save the simulation snapshot.
10. Create a planned cross and observation protocol.

## Simulation result page

Always display:

- result authority label;
- input parent identities and direction;
- genotype evidence statuses;
- model and catalog release versions;
- exact versus sampled calculation;
- assumptions and exclusions;
- probability distributions;
- target probability and population calculation;
- scientific sources and applicable populations;
- warnings and abstentions;
- reproducibility identifier;
- action to create a breeding plan.

## Material identity workflow

Support printable and QR labels for seed lots, plants, pollinated flowers, fruits, tissue samples, and seed harvests. Scanning must open the correct record and prevent operators from accidentally assigning observations to the wrong biological material.

## Phenotype capture workflow

1. Scan or select the material.
2. Choose the observation protocol.
3. Show capture instructions and calibration requirements.
4. Capture/upload images.
5. Run quality checks before analysis.
6. Present deterministic measurements and model suggestions.
7. Require human correction/approval for authoritative observations.
8. Store raw image, derived mask, calibration, model version, edits, and final observation.

## Scientific review workflow

Candidate claims move through:

```text
draft -> extracted -> curator_review -> independent_review -> approved -> released
                                           |                |
                                           -> rejected      -> superseded
```

The interface must show source passages next to extracted assertions. Reviewers must be able to record disagreement, limitations, species/population scope, and prohibited interpretations.

## AI Copilot UX

AI responses must show:

- answer class: explanation, evidence summary, workflow help, or tool result;
- citations to source passages or internal records;
- tools called and simulation run identifiers;
- uncertainty and missing data;
- actions that require human approval;
- an obvious way to report an incorrect answer.

Do not display AI prose as an authoritative genotype or scientific rule.
''')

write(ROOT / 'docs/03_DATA_MODEL_AND_MIGRATIONS.md', r'''
# Data Model and Migration Specification

## Design rules

- Use UUIDv7 or another sortable globally unique identifier consistently.
- Every workspace-owned authoritative table carries `workspace_id`.
- Every mutable authoritative table carries creation/update actor and timestamps.
- Scientific releases and simulation snapshots are immutable after publication/completion.
- Soft deletion is allowed only where legal/audit requirements permit it; biological lineage records should generally be retired, not erased.
- References to biological material must identify the exact accession, seed lot, individual plant, sample, or harvest—not only a cultivar string.
- Units use controlled terms and canonical conversions.

## Major schema domains

### Identity and authorization

`users`, `workspaces`, `workspace_memberships`, `roles`, `permissions`, `api_keys`, `sessions`, `audit_events`.

### Scientific catalog

`catalog_releases`, `loci`, `locus_aliases`, `alleles`, `allele_aliases`, `sequence_variants`, `reference_assemblies`, `markers`, `marker_assays`, `inheritance_models`, `model_rules`, `locus_interactions`, `phenotype_terms`, `evidence_assertions`, `evidence_applicability`, `publications`, `source_passages`, `scientific_reviews`, `catalog_change_requests`.

### Germplasm and material identity

`taxa`, `germplasm_accessions`, `germplasm_names`, `source_organizations`, `acquisitions`, `seed_lots`, `seed_packets`, `plants`, `plant_locations`, `material_labels`, `material_events`, `fruits`, `tissue_samples`, `seed_harvests`.

### Breeding and pedigree

`crosses`, `cross_parents`, `pollination_events`, `cross_verifications`, `progeny_families`, `generation_records`, `pedigree_edges`, `selection_goals`, `selection_criteria`, `selection_decisions`.

### Genotyping

`genotype_assays`, `genotype_calls`, `genotype_call_evidence`, `haplotypes`, `haplotype_calls`, `variant_sets`, `call_sets`, `quality_metrics`.

### Experiments and observations

`investigations`, `studies`, `environments`, `locations`, `experimental_designs`, `blocks`, `plots`, `observation_units`, `observation_variables`, `traits`, `methods`, `scales`, `units`, `observation_sessions`, `observations`, `environment_observations`, `treatments`, `devices`, `calibrations`, `quality_control_events`.

### Media and vision

`media_assets`, `image_captures`, `capture_protocols`, `image_quality_results`, `annotations`, `annotation_revisions`, `segmentation_masks`, `derived_measurements`, `vision_models`, `vision_model_versions`, `vision_evaluations`.

### Simulation and statistics

`simulation_models`, `simulation_model_versions`, `simulation_runs`, `simulation_inputs`, `simulation_results`, `simulation_artifacts`, `selection_plans`, `model_training_datasets`, `statistical_models`, `statistical_model_versions`, `model_cards`, `model_evaluations`, `model_promotions`.

### Research and AI

`research_documents`, `document_versions`, `document_chunks`, `embedding_records`, `candidate_claims`, `ai_runs`, `ai_messages`, `ai_tool_calls`, `ai_evaluations`, `human_feedback`.

### Jobs and operations

`jobs`, `job_attempts`, `job_logs`, `outbox_events`, `idempotency_records`, `backup_records`, `restore_tests`.

## Critical constraints

- Exactly one maternal and one paternal role for an ordinary biparental cross; support alternative reproduction types explicitly.
- A plant cannot be its own ancestor and pedigree cycles must be rejected.
- An offspring biological material has no more than one origin event.
- A verified genotype call requires an approved assay/evidence method and cannot be created by AI.
- An assumed or inferred call cannot masquerade as verified.
- Published catalog releases cannot be mutated.
- Completed simulation input and result payloads cannot be mutated.
- Observation values must match the declared scale and unit.
- Image-derived authoritative measurements require a passing capture-quality state and approved human or model pathway.
- Model promotion requires an evaluation record and approval.

## Catalog import

The supplied workbook is an import source, not the schema. Implement a staged importer:

1. Load workbook/CSV into a staging schema.
2. Validate identifiers, source references, evidence grades, and required fields.
3. Generate normalized draft records.
4. Produce a reconciliation report showing information preserved, transformed, rejected, or unresolved.
5. Require curator approval before a catalog release is published.
6. Preserve original row payloads and file hash for provenance.

## Migration policy

- Forward-only production migrations with tested backups and explicit forward-repair procedures.
- Every destructive change requires a data migration, verification query, and rollback/restore decision.
- Seed data is versioned and idempotent.
- Migration tests run against an empty database and a representative prior schema snapshot.
- No application boot-time schema mutation.
''')

write(ROOT / 'docs/04_SIMULATION_ENGINE_SPEC.md', r'''
# Simulation Engine Specification

## Authority model

Simulation results are authoritative only within the approved model, validated applicability, and declared inputs. The engine must abstain when a requested claim is unsupported.

## Shared run contract

Every run stores:

- run ID and workspace;
- parent/material identifiers;
- complete normalized input;
- genotype evidence states;
- catalog release;
- model type and version;
- exact or sampled mode;
- random seed when sampled;
- result distribution;
- assumptions, warnings, exclusions, and abstentions;
- code/worker version;
- timestamps and content hashes.

## Engine A — Exact nuclear inheritance

Implement:

- diploid single-locus crosses;
- multiple alleles;
- independent multi-locus crosses;
- complete, incomplete, and codominance through declarative phenotype rules;
- penetrance and expressivity distributions where explicitly supported;
- lethal combinations with conception and surviving-offspring distributions kept separate;
- exact rational probability arithmetic.

Never infer dominance from capitalization.

## Engine B — Linked loci

Inputs include chromosome, haplotype phase, and recombination fraction or validated map distance. Support coupling and repulsion. Unknown phase must generate scenario-separated or marginalized results rather than a hidden assumption.

Invariants:

- `r = 0` yields parental haplotypes only;
- `r = 0.5` converges to independent assortment;
- gamete probabilities sum exactly to 1.

## Engine C — Cytoplasmic and maternal models

Keep parent direction. Support maternal cytoplasm, nuclear restorer loci, maternal effects, and reciprocal-cross outcome differences through explicit model rules.

## Engine D — Epistasis and rule graphs

Use versioned declarative rules over normalized genotypes and context. Rules require priorities, conflict detection, source assertions, applicability, and an abstention result when no approved rule resolves the request.

## Engine E — Host–pathogen interaction

Inputs include host allele/genotype, pathogen taxon, strain/isolate/race/pathotype, effector status, environment, and evidence applicability. Results are context-specific and may be `supported_resistant`, `supported_susceptible`, `conflicting`, or `insufficient_evidence`.

## Engine F — Uncertain genotypes

Represent parent genotype as a weighted distribution. Marginalize across parent distributions and genetic outcomes:

```text
P(offspring) = sum P(parentA genotype) * P(parentB genotype) * P(offspring | parent genotypes)
```

Expose how much uncertainty comes from parent identity, genotype inference, or the biological model.

## Engine G — Target recovery

For independent target probability `p` and desired confidence `c`:

```text
N = ceil(log(1-c) / log(1-p))
```

Apply germination, survival, assay success, and phenotyping sensitivity as separate transparent operational factors. Never label the result a guarantee.

## Engine H — Monte Carlo fallback

Use only when exact state space exceeds configured thresholds. Store random seed, sample count, convergence diagnostics, confidence intervals, and exact-fallback reason. Provide deterministic repeatability from the stored seed.

## Engine I — Forward breeding-program simulation

Use an isolated R worker and AlphaSimR or a justified alternative. Support founder population creation, mating design, selection cycles, genomic selection experiments, and comparison of strategies. These stochastic research simulations must never overwrite exact calculator results.

## Engine J — Quantitative and genomic prediction

Implement framework and registry support for:

- mixed models and BLUP;
- GBLUP/genomic selection;
- Bayesian models when justified;
- multi-trait selection;
- genotype-by-environment models;
- reaction norms;
- selection indices and Pareto optimization.

Promotion requires frozen training data, leakage-resistant validation, uncertainty, calibration, model cards, and applicability boundaries. Without adequate data, return `model_unavailable`.

## Engine K — In-silico genotype manipulation

Allow users to explore hypothetical allele edits or selections as scenarios. Label outputs as virtual scenarios. Do not provide laboratory gene-editing instructions or imply that an edit will produce a phenotype outside an approved model.

## Performance targets

- Typical exact crosses: interactive response under 500 ms server compute.
- Larger exact crosses: under 2 seconds or move to a job with progress.
- No unbounded Cartesian enumeration.
- Sparse weighted distributions and dynamic programming where applicable.
- Configured hard limits and cost estimates before expensive simulations.

## Testing

Use golden fixtures, property-based tests, an independent reference implementation, randomized cross comparisons, and mathematical invariants. Every model version must be reproducible after future catalog releases.
''')

write(ROOT / 'docs/05_AI_AND_RAG_SPEC.md', r'''
# AI and Research Retrieval Specification

## Principle

AI assists scientific work; it does not define biological truth. Deterministic and validated statistical engines produce calculations. Humans approve scientific claims and verified genotypes.

## AI use cases

- ingest public papers and user-provided research;
- extract candidate loci, alleles, variants, populations, methods, outcomes, and limitations into draft schemas;
- identify nomenclature aliases and conflicts;
- compare supporting and conflicting studies;
- explain simulations and uncertainty in plain language;
- help users plan observations and breeding workflows;
- convert dictated or typed notes into draft observations;
- explain image-analysis results and request human corrections;
- query the user's breeding records through authorized tools.

## Prohibited AI actions

- activating catalog rules;
- marking genotype calls verified;
- inventing citations, alleles, effect sizes, or inheritance models;
- calculating authoritative probabilities in prose instead of calling simulation tools;
- diagnosing disease conclusively from an image;
- bypassing permissions or revealing records from another workspace;
- training on private user data without an explicit governed workflow.

## Retrieval architecture

Use hybrid retrieval over:

- normalized publications and evidence assertions;
- exact source passages with page/section metadata;
- catalog aliases and nomenclature;
- user-authorized breeding records;
- model cards and system documentation.

Embeddings are a discovery index. The source of truth is the stored source passage and normalized record. Every cited answer must be resolvable to source text.

## Research ingestion pipeline

```text
file/identifier intake
-> malware/type validation
-> content hashing and versioning
-> parser/extractor
-> page and section segmentation
-> metadata resolution
-> candidate structured extraction
-> source-passage linking
-> nomenclature matching
-> conflict and duplicate detection
-> curator review
-> independent review
-> approved release
```

OCR is a fallback. Preserve the original file and extraction quality.

## Tool contracts

The bounded agent may call tools such as:

- `search_evidence`
- `read_source_passage`
- `get_catalog_entity`
- `compare_evidence`
- `run_exact_simulation`
- `run_advanced_simulation`
- `calculate_population_target`
- `get_material_record`
- `create_draft_observation`
- `create_draft_candidate_claim`
- `explain_vision_result`

Write tools require explicit permissions and human confirmation. All tool calls are recorded.

## AI answer schema

Every answer contains:

- answer text;
- answer class;
- citations;
- internal record references;
- tool calls and run IDs;
- confidence/uncertainty statement;
- missing-data list;
- prohibited inference warnings;
- actions requiring approval.

## Provider strategy

Use AI SDK 6 with provider abstraction. Support at least one hosted provider and a deterministic/mock provider for tests. Provider failures must degrade gracefully without breaking core breeding or simulation features.

## Evaluation

Maintain versioned eval suites for:

- citation correctness;
- unsupported claim rate;
- tool selection;
- schema adherence;
- nomenclature accuracy;
- answer completeness;
- cross-workspace data isolation;
- prompt-injection resistance;
- abstention behavior;
- regression across model/provider upgrades.

No AI feature is promoted solely because sample answers look persuasive.
''')

write(ROOT / 'docs/06_PHENOTYPE_VISION_SPEC.md', r'''
# Phenotype Vision Specification

## Goal

Provide standardized image capture, deterministic measurements, annotation workflows, and validated task-specific models for pepper plants and fruit.

## Capture protocols

Each protocol defines:

- material type and developmental stage;
- required views;
- background and lighting;
- distance and camera orientation;
- scale object;
- color calibration target;
- device metadata;
- file format and resolution;
- operator instructions;
- allowable quality thresholds.

Supported initial protocols:

1. Detached fruit on calibrated background.
2. Whole plant, front and side views.
3. Leaf close-up.
4. Fruit cross-section.
5. Symptom close-up with context image.

## Quality gate

Before analysis, assess:

- blur/sharpness;
- exposure and clipping;
- occlusion;
- required calibration visibility;
- background compliance;
- object completeness;
- resolution;
- duplicate or near-duplicate capture;
- metadata consistency.

A failed capture does not silently produce authoritative measurements.

## Deterministic pipeline

Use OpenCV/PlantCV-style processing for:

- color correction;
- geometric calibration;
- object segmentation where reliable;
- length, width, area, perimeter, aspect ratio, curvature, and shape descriptors;
- color distributions in calibrated spaces;
- count and location measurements;
- mask and overlay generation.

Every measurement records protocol, algorithm version, calibration, raw image, mask, and human edits.

## Learned models

Task-specific models may cover:

- fruit detection and instance segmentation;
- maturity-stage classification;
- fruit shape classes;
- anthocyanin patterns;
- leaf or fruit symptom candidate classification;
- plant architecture measurements;
- flower/fruit counting.

SAM-style models may accelerate annotation but their zero-shot masks are not automatically ground truth.

## Human-in-the-loop

- Users can correct masks, object counts, maturity stage, and classification.
- Corrections are stored as new annotation revisions.
- Model outputs and human observations remain distinguishable.
- Active-learning queues prioritize uncertain and out-of-distribution examples.

## Dataset governance

Split datasets by accession/family, environment/season, and capture device where possible. Prevent near-duplicate leakage. Track labeler, protocol, label version, consent/license, and exclusion reasons.

## Evaluation

Depending on task, report:

- IoU/Dice and boundary error;
- mean absolute measurement error against physical measurements;
- precision, recall, F1, calibration, and confusion matrix;
- device-, accession-, color-, size-, and environment-stratified performance;
- out-of-distribution detection;
- human correction rate.

An image-analysis feature must state the validated protocol and applicability.
''')

write(ROOT / 'docs/07_SCIENTIFIC_GOVERNANCE.md', r'''
# Scientific Governance

## Evidence states

```text
candidate -> extracted -> curator_review -> independent_review -> approved -> released
                                  |                   |
                                  -> rejected         -> superseded/retired
```

## Evidence grades

- **A:** functionally validated causal relationship with strong replicated support.
- **B:** replicated mapping/association or strong mechanistic support.
- **C:** population-specific QTL or association.
- **D:** candidate, preliminary, conflicting, or weakly replicated.
- **U:** unresolved or insufficient.

Evidence grade and executable prediction eligibility are separate decisions.

## Required assertion fields

- narrow claim text;
- claim type;
- species and population;
- biological material/accessions;
- allele or variant identity;
- inheritance/model interpretation;
- experimental method;
- environmental conditions;
- supporting passages;
- limitations and exclusions;
- conflicting evidence;
- evidence grade;
- reviewer decisions;
- release version.

## Genotype evidence states

- verified by sequencing;
- verified by approved marker assay;
- imported verified source;
- pedigree inferred;
- phenotype inferred;
- user assumed;
- unknown;
- conflicting.

Only approved verification pathways can create a verified state.

## Publication governance

- Store DOI/PMID/other identifiers and source version.
- Do not redistribute copyrighted full text unless permitted.
- Store legal source links and user-supplied files according to access rights.
- Track retractions, corrections, and superseding evidence.

## Rule activation

An executable model rule requires:

- approved evidence assertions;
- normalized alleles/variants;
- explicit applicability;
- declarative rule and test fixtures;
- scientific reviewer approval;
- software reviewer approval;
- catalog release publication.

## Change control

Published releases are immutable. Corrections create a new release. Historical simulations retain the prior catalog/model version and remain reproducible.
''')

write(ROOT / 'docs/08_API_AND_CONTRACTS.md', r'''
# API and Contract Strategy

## Internal approach

Use typed application services for server-side Next.js workflows and versioned HTTP/job contracts at runtime boundaries. Do not expose database tables directly as APIs.

## Core resource families

- identity/workspaces;
- germplasm and seed lots;
- plants and material events;
- crosses, pollination, harvests, progeny, and pedigrees;
- genotypes and assays;
- catalog loci, alleles, claims, sources, reviews, and releases;
- studies, observation variables, sessions, and observations;
- images, annotations, and measurements;
- simulation models, runs, results, and selection plans;
- research documents and AI runs;
- jobs, models, audit, exports, and backups.

## API rules

- Validate every input at the boundary.
- Authorize every operation against workspace and action.
- Use idempotency keys for create/mutate operations vulnerable to retries or repeated clicks.
- Use optimistic concurrency or version checks for user-edited records.
- Return stable error codes with operator-safe messages.
- Use cursor pagination.
- Do not expose worker internal paths, secrets, model credentials, or raw stack traces.
- Generate OpenAPI for public/versioned HTTP APIs.

## Interoperability

Implement selected BrAPI 2.1 endpoints only after internal domain contracts stabilize. Prioritize germplasm, seed lots, crosses, pedigree, studies, observation variables, observations, images, samples, and variants. Provide MIAPPE 1.2 and MCPD-compatible exports.

## Job contract

All worker requests and results use JSON Schema under `contracts/`, include schema versions, and are validated by both producer and consumer. Large files use signed object-store references and content hashes rather than inline base64.
''')

write(ROOT / 'docs/09_SECURITY_PRIVACY_OPERATIONS.md', r'''
# Security, Privacy, and Operations

## Roles

Minimum roles:

- workspace owner;
- breeder;
- technician/observer;
- scientific curator;
- independent scientific reviewer;
- data analyst;
- read-only collaborator;
- system administrator.

## Separation of duties

- Claim extractor cannot self-approve a claim.
- Curator approval and independent review are separately attributable.
- AI cannot approve anything.
- Model trainer cannot promote a model without required evaluation and approval.
- Destructive administrative actions require elevated permission and audit.

## Security controls

- secure session cookies and CSRF protection;
- rate limiting and abuse controls;
- strict file validation and malware scanning path;
- signed object-store access;
- encrypted transport and encrypted backups;
- secret manager or environment-injected secrets, never database-stored plaintext;
- least-privilege database and object-store credentials;
- dependency and container scanning;
- security headers and CSP;
- complete audit trail for authoritative changes;
- cross-workspace isolation tests;
- AI prompt-injection and data-exfiltration defenses.

## Reliability controls

- idempotent mutations and jobs;
- transactional outbox for side effects;
- bounded retries with dead-letter/manual recovery;
- health/readiness checks;
- structured logs with correlation IDs;
- metrics for web, database, queue, workers, AI, vision, and simulations;
- trace context across job boundaries;
- backup schedules and restore drills;
- object integrity hashes;
- graceful degradation when AI or workers are unavailable.

## Data ownership

Users can export their breeding, phenotype, genotype, media metadata, and model-run records in documented nonproprietary formats. Account deletion must respect scientific audit and collaboration requirements and clearly distinguish deletion, anonymization, and retained shared records.
''')

write(ROOT / 'docs/10_TEST_AND_VALIDATION_PLAN.md', r'''
# Test and Validation Plan

## Test layers

### Static quality

- formatting;
- linting;
- strict TypeScript;
- Python type checking and lint;
- R package/script checks;
- schema validation;
- dependency and secret scanning.

### Exact genetics

- golden crosses;
- property-based testing;
- probability sum and nonnegativity;
- genotype normalization invariance;
- parent-swap invariance for nuclear models;
- parent-direction differences for maternal models;
- linkage boundary tests;
- lethal/conception/survival separation;
- exact-versus-Monte-Carlo convergence;
- independent reference implementation parity.

### Database/domain

- constraints and transaction tests;
- migration tests;
- pedigree cycle prevention;
- authorization and workspace isolation;
- idempotency and concurrency;
- immutable release/run behavior;
- import reconciliation.

### Web application

- component tests;
- server action/API integration tests;
- Playwright end-to-end workflows;
- keyboard and screen-reader checks;
- automated accessibility scanning;
- responsive visual regression;
- slow/error/degraded-state tests;
- repeated-click and retry behavior.

### AI

- tool schema tests;
- mock-provider deterministic tests;
- citation-grounding evals;
- prompt-injection tests;
- cross-workspace leakage tests;
- abstention and unsupported-claim tests;
- provider-failure degradation.

### Vision

- image-quality fixtures;
- calibration tests;
- measurement comparison with physical ground truth;
- segmentation and classifier benchmarks;
- leakage audits;
- device/accession/environment holdouts;
- human correction workflows.

### Operations

- Docker Compose smoke test;
- clean installation;
- backup and restore drill;
- worker restart/retry behavior;
- migration failure behavior;
- observability and alert test;
- export/import round trip.

## Completion rule

No workstream may report completion only from unit tests. User-visible capabilities require integrated and end-to-end proof. Scientific prediction capabilities require model-specific validation and applicability evidence.
''')

write(ROOT / 'docs/11_DELIVERY_AND_RELEASE_PLAN.md', r'''
# Delivery and Release Plan

## Release slices

### Slice 1 — Foundation and catalog governance

Monorepo, web shell, identity, workspaces, roles, PostgreSQL, object storage, jobs, audit, catalog importer, review states, and seeded catalog.

### Slice 2 — Exact genetics and simulation UX

Exact rational engine, independent loci, multi-allelic rules, uncertain genotypes, target recovery, immutable run storage, result UI, and tests.

### Slice 3 — Breeding ledger

Germplasm, seed lots, plants, labels, crosses, pollination, harvests, progeny, pedigree, selection goals, observations, and exports.

### Slice 4 — Advanced genetics

Linkage, maternal/cytoplasmic inheritance, epistasis, host–pathogen contexts, Monte Carlo fallback, and advanced model UI.

### Slice 5 — AI research and copilot

Document ingestion, hybrid retrieval, candidate extraction, review queue, simulation tools, AI UI, evals, and provider degradation.

### Slice 6 — Phenotype laboratory

Capture protocols, quality gates, deterministic measurements, annotations, corrections, model registry, and vision job UI.

### Slice 7 — Research simulation and model framework

R/Python model workers, AlphaSimR job templates, training dataset registry, model cards, evaluation and promotion gates. Quantitative predictions remain disabled until data gates pass.

### Slice 8 — Interoperability and hardening

Selected BrAPI endpoints, MIAPPE/MCPD exports, performance, security, restore drill, accessibility audit, deployment docs, and release evidence.

## Merge order

1. Platform contracts and repository baseline.
2. Database schema and catalog import.
3. Exact engine.
4. Core web workflows.
5. Breeding ledger.
6. Advanced simulation contracts.
7. AI and workers.
8. Vision and model registry.
9. Integrated validation and release hardening.

## Release strategy

Use feature flags for incomplete research capabilities. Seed a demonstration workspace. Run migrations and backup before deployment. Maintain the previous deployable image and catalog/model versions for rollback. Prefer forward repair for schema changes after data migration.
''')

write(ROOT / 'docs/12_DEFINITION_OF_DONE.md', r'''
# Definition of Done

The application is complete for this handoff only when all of the following are true.

## Product

- The web interface is functional across the required route map.
- A user can complete onboarding and create or use a seeded workspace.
- A full breeding cycle can be recorded without an external spreadsheet.
- A supported cross can be simulated and converted into a breeding plan.
- Evidence and uncertainty are visible at the point of prediction.
- Unsupported predictions abstain clearly.

## Science

- The supplied catalog imports without silent information loss.
- No spreadsheet row is directly treated as executable logic.
- All executable rules are versioned, source-linked, applicability-limited, and tested.
- Material identity, parent direction, genotype evidence state, and provenance are preserved.

## Engineering

- Repository package boundaries are enforced.
- Database migrations, seed import, and rollback/repair paths are documented.
- Exact simulation is deterministic and independently verified.
- Jobs are idempotent, observable, retryable, and bounded.
- AI and workers can be unavailable without breaking core workflows.
- Security, authorization, and workspace isolation tests pass.
- CI runs formatting, lint, types, tests, builds, security checks, and E2E smoke.

## UX

- WCAG 2.2 AA blocking issues are resolved.
- Mobile/tablet/desktop layouts are tested.
- Empty, loading, error, partial, degraded, and permission states exist.
- High-risk actions have confirmation and recovery behavior.
- Nontechnical users receive actionable guidance rather than raw errors.

## Operations

- Docker Compose starts the supported local stack.
- Environment variables are documented and validated.
- Backup and restore are proven.
- Logs, metrics, and health checks are present.
- All user data can be exported.

## Reporting

The final agent report maps every acceptance criterion to evidence, lists all files changed, reports exact commands and outcomes, classifies failures, and identifies residual scientific gates without overstating readiness.
''')

write(ROOT / 'docs/13_ENVIRONMENT_AND_DEPLOYMENT.md', r'''
# Environment and Deployment

## Supported local profile

- Linux, macOS, or Windows with WSL2.
- Docker Engine/Compose.
- Node.js 24 LTS.
- pnpm with Corepack.
- Optional local Python/R tooling for worker development; containers remain authoritative.

## Required services

```text
web
postgres:18
object-storage (MinIO-compatible locally)
worker-python
worker-r
backup
optional local mail catcher
optional observability stack
```

## Environment groups

- application URLs and environment;
- database connection and pool controls;
- auth/session secrets and trusted origins;
- object storage endpoint/bucket/credentials;
- queue controls;
- AI provider keys and selected model IDs;
- embedding provider/model;
- worker endpoints and shared authentication;
- telemetry endpoints;
- upload and job limits;
- feature flags.

Provide `.env.example` with no real secrets and a typed configuration validator that fails early with actionable messages.

## Deployment requirements

- reproducible container builds;
- non-root runtime users;
- health/readiness checks;
- migration job separated from app startup;
- persistent PostgreSQL and object-storage volumes;
- TLS at reverse proxy;
- image and dependency provenance where practical;
- deployment and rollback runbooks;
- catalog and model version rollback independent from application rollback.
''')

write(ROOT / 'docs/14_REPO_STRUCTURE_AND_CONVENTIONS.md', r'''
# Repository Structure and Conventions

## Monorepo

Use pnpm workspaces and Turborepo or an equivalent task graph. Keep app deployables and domain packages explicit.

## TypeScript rules

- strict mode;
- no `any` at domain boundaries;
- branded identifiers or validated ID types;
- discriminated unions for scientific states;
- exhaustive switches;
- no floating-point probability authority in the exact engine;
- schema validation for untrusted data;
- domain errors mapped to stable user-facing error codes.

## Database rules

- repositories/services own transactions;
- no direct database access from UI components;
- no generic unscoped CRUD for workspace data;
- migrations are reviewed artifacts;
- generated clients/types are regenerated from source contracts;
- query performance is tested for pedigree, catalog search, and observation pages.

## UI rules

- server components by default;
- client components only for interactive state;
- URL-addressable filters and views;
- accessible shared form primitives;
- consistent status badges and authority labels;
- prevent duplicate submissions at both UI and server layers;
- optimistic UI only when rollback is safe and clear.

## Testing rules

- test behavior, invariants, and contracts;
- no snapshot-only scientific tests;
- fixtures state their scientific source or synthetic purpose;
- flaky tests are defects;
- tests must run in CI without private provider credentials.
''')

write(ROOT / 'docs/15_VERIFIED_TECH_AND_STANDARDS_BASELINE.md', r'''
# Verified Technology and Standards Baseline

Verified on 2026-07-22. The implementation agent must still select current compatible patch versions and record them in the final report.

| Technology/standard | Baseline | Official source |
|---|---|---|
| Node.js | 24 LTS; production should use an LTS line | https://nodejs.org/en/about/previous-releases |
| Next.js | 16.2 major/minor baseline | https://nextjs.org/blog/next-16-2 |
| AI SDK | 6.0 generation | https://ai-sdk.dev/docs/migration-guides/migration-guide-6-0 |
| PostgreSQL | 18 current stable major; use current minor | https://www.postgresql.org/docs/current/index.htm |
| MIAPPE | 1.2, October 2024; compatible with 1.1 | https://www.miappe.org/releases/ |
| BrAPI | 2.1 latest stable and recommended for new development | https://brapi.org/specification |

Do not use PostgreSQL 19 beta in production. Do not bind the application to every BrAPI endpoint; implement only approved modules after internal contracts stabilize.
''')

# JSON schemas / contracts
schema_base = 'https://json-schema.org/draft/2020-12/schema'

dump_json(ROOT / 'contracts' / 'simulation-request.schema.json', {
    '$schema': schema_base,
    '$id': 'https://capsicum.local/schemas/simulation-request.schema.json',
    'title': 'SimulationRequest',
    'type': 'object',
    'additionalProperties': False,
    'required': ['schemaVersion', 'workspaceId', 'modelType', 'catalogReleaseId', 'parents', 'loci'],
    'properties': {
        'schemaVersion': {'const': '1.0'},
        'workspaceId': {'type': 'string', 'minLength': 1},
        'modelType': {'enum': ['exact_nuclear', 'linked', 'maternal', 'epistasis', 'host_pathogen', 'uncertain', 'monte_carlo', 'forward_program', 'quantitative', 'genomic', 'gxe', 'virtual_scenario']},
        'catalogReleaseId': {'type': 'string', 'minLength': 1},
        'parents': {
            'type': 'object', 'additionalProperties': False,
            'required': ['maternal', 'paternal'],
            'properties': {
                'maternal': {'$ref': '#/$defs/parent'},
                'paternal': {'$ref': '#/$defs/parent'},
            }
        },
        'loci': {'type': 'array', 'minItems': 1, 'items': {'type': 'string'}},
        'target': {'type': ['object', 'null']},
        'operationalFactors': {'type': ['object', 'null']},
        'randomSeed': {'type': ['integer', 'null']},
        'sampleCount': {'type': ['integer', 'null'], 'minimum': 1},
        'context': {'type': 'object'},
    },
    '$defs': {
        'parent': {
            'type': 'object', 'additionalProperties': False,
            'required': ['materialId', 'genotypeStates'],
            'properties': {
                'materialId': {'type': 'string', 'minLength': 1},
                'genotypeStates': {'type': 'array', 'items': {'$ref': '#/$defs/genotypeState'}},
            }
        },
        'genotypeState': {
            'type': 'object', 'additionalProperties': False,
            'required': ['locusId', 'status', 'possibilities'],
            'properties': {
                'locusId': {'type': 'string'},
                'status': {'enum': ['verified', 'inferred', 'assumed', 'unknown', 'conflicting']},
                'possibilities': {
                    'type': 'array', 'minItems': 1,
                    'items': {
                        'type': 'object', 'additionalProperties': False,
                        'required': ['alleles', 'probability'],
                        'properties': {
                            'alleles': {'type': 'array', 'minItems': 2, 'maxItems': 2, 'items': {'type': 'string'}},
                            'probability': {'type': 'object', 'required': ['numerator', 'denominator'], 'properties': {'numerator': {'type': 'integer', 'minimum': 0}, 'denominator': {'type': 'integer', 'minimum': 1}}},
                        }
                    }
                }
            }
        }
    }
})

dump_json(ROOT / 'contracts' / 'simulation-result.schema.json', {
    '$schema': schema_base,
    '$id': 'https://capsicum.local/schemas/simulation-result.schema.json',
    'title': 'SimulationResult',
    'type': 'object',
    'additionalProperties': False,
    'required': ['schemaVersion', 'runId', 'authority', 'modelVersion', 'catalogReleaseId', 'calculationMode', 'distributions', 'assumptions', 'warnings', 'abstentions', 'contentHash'],
    'properties': {
        'schemaVersion': {'const': '1.0'},
        'runId': {'type': 'string'},
        'authority': {'enum': ['exact_supported', 'conditional_supported', 'research_only', 'model_derived', 'unsupported']},
        'modelVersion': {'type': 'string'},
        'catalogReleaseId': {'type': 'string'},
        'calculationMode': {'enum': ['exact', 'sampled', 'statistical']},
        'distributions': {'type': 'array', 'items': {'type': 'object'}},
        'targetRecovery': {'type': ['object', 'null']},
        'assumptions': {'type': 'array', 'items': {'type': 'string'}},
        'warnings': {'type': 'array', 'items': {'type': 'string'}},
        'abstentions': {'type': 'array', 'items': {'type': 'string'}},
        'sourceAssertionIds': {'type': 'array', 'items': {'type': 'string'}},
        'randomSeed': {'type': ['integer', 'null']},
        'sampleCount': {'type': ['integer', 'null']},
        'contentHash': {'type': 'string', 'minLength': 32},
    }
})

dump_json(ROOT / 'contracts' / 'ai-answer.schema.json', {
    '$schema': schema_base,
    '$id': 'https://capsicum.local/schemas/ai-answer.schema.json',
    'title': 'AiAnswer',
    'type': 'object',
    'additionalProperties': False,
    'required': ['schemaVersion', 'answerClass', 'text', 'citations', 'toolCalls', 'uncertainty', 'missingData', 'requiresApproval'],
    'properties': {
        'schemaVersion': {'const': '1.0'},
        'answerClass': {'enum': ['workflow_help', 'evidence_summary', 'simulation_explanation', 'record_summary', 'draft_observation', 'candidate_claim', 'abstention']},
        'text': {'type': 'string'},
        'citations': {'type': 'array', 'items': {'type': 'object', 'required': ['sourceId', 'passageId'], 'properties': {'sourceId': {'type': 'string'}, 'passageId': {'type': 'string'}, 'locator': {'type': ['string', 'null']}}}},
        'toolCalls': {'type': 'array', 'items': {'type': 'object'}},
        'uncertainty': {'type': 'string'},
        'missingData': {'type': 'array', 'items': {'type': 'string'}},
        'warnings': {'type': 'array', 'items': {'type': 'string'}},
        'requiresApproval': {'type': 'boolean'},
    }
})

dump_json(ROOT / 'contracts' / 'phenotype-observation.schema.json', {
    '$schema': schema_base,
    '$id': 'https://capsicum.local/schemas/phenotype-observation.schema.json',
    'title': 'PhenotypeObservation',
    'type': 'object',
    'additionalProperties': False,
    'required': ['schemaVersion', 'workspaceId', 'observationUnitId', 'observationVariableId', 'value', 'status', 'provenance'],
    'properties': {
        'schemaVersion': {'const': '1.0'},
        'workspaceId': {'type': 'string'},
        'observationUnitId': {'type': 'string'},
        'observationVariableId': {'type': 'string'},
        'value': {},
        'unitId': {'type': ['string', 'null']},
        'status': {'enum': ['draft', 'observed', 'machine_suggested', 'human_corrected', 'reviewed', 'rejected']},
        'provenance': {
            'type': 'object', 'required': ['method', 'actorType'],
            'properties': {
                'method': {'type': 'string'},
                'actorType': {'enum': ['human', 'deterministic_algorithm', 'machine_learning_model', 'ai_assistant', 'import']},
                'actorId': {'type': ['string', 'null']},
                'modelVersionId': {'type': ['string', 'null']},
                'imageCaptureId': {'type': ['string', 'null']},
                'protocolId': {'type': ['string', 'null']},
            }
        },
        'qualityFlags': {'type': 'array', 'items': {'type': 'string'}},
    }
})

dump_json(ROOT / 'contracts' / 'catalog-import-record.schema.json', {
    '$schema': schema_base,
    '$id': 'https://capsicum.local/schemas/catalog-import-record.schema.json',
    'title': 'CatalogImportRecord',
    'type': 'object',
    'required': ['sourceFileHash', 'sourceSheet', 'sourceRow', 'rawPayload', 'normalizationStatus'],
    'properties': {
        'sourceFileHash': {'type': 'string'},
        'sourceSheet': {'type': 'string'},
        'sourceRow': {'type': 'integer', 'minimum': 1},
        'rawPayload': {'type': 'object'},
        'normalizationStatus': {'enum': ['pending', 'normalized', 'requires_review', 'rejected']},
        'normalizedEntityIds': {'type': 'array', 'items': {'type': 'string'}},
        'issues': {'type': 'array', 'items': {'type': 'object'}},
    }
})

write(ROOT / 'contracts' / 'README.md', r'''
# Machine-Readable Contracts

These schemas define minimum cross-runtime contracts. The implementation agent may extend them through additive, versioned changes but must not silently weaken required provenance, authority, uncertainty, or identity fields.

Generate TypeScript, Python, and R validation/types from one canonical contract source or maintain contract tests that prove equivalent validation across runtimes.
''')

write(ROOT / 'starter-repo-blueprint' / 'REPOSITORY_TREE.md', r'''
# Expected Repository Tree

```text
capsicum-breeding-intelligence/
  AGENTS.md
  README.md
  package.json
  pnpm-workspace.yaml
  turbo.json
  tsconfig.base.json
  .env.example
  .gitignore
  .github/workflows/ci.yml
  apps/
    web/
    worker-python/
    worker-r/
  packages/
    auth/
    database/
    contracts/
    config/
    genetics-core/
    genetics-advanced/
    scientific-catalog/
    breeding-domain/
    observation-domain/
    simulation-domain/
    ai/
    vision/
    jobs/
    storage/
    observability/
    ui/
    test-utils/
  infra/
    compose/
    docker/
    backup/
    observability/
  docs/
  scripts/
```

The coding agent may refine names but must preserve domain boundaries and dependency direction.
''')

write(ROOT / 'starter-repo-blueprint' / 'AGENTS.md', r'''
# Repository Agent Rules

- Read the project handoff pack before making architecture changes.
- Scientific claims, verified genotype states, catalog releases, and promoted models require explicit governed workflows.
- The exact genetics engine is pure, deterministic, and free from database, network, React, and AI dependencies.
- AI cannot create authoritative science or verified genotype data.
- Do not infer dominance from allele capitalization.
- Do not silently replace unknown inputs with defaults.
- Preserve maternal and paternal direction.
- All workspace data access is explicitly scoped and authorized.
- All retryable mutations and jobs are idempotent.
- Generated files are changed through their source generator.
- No task is complete without tests and evidence.
- Unsupported prediction requests must abstain.
''')

write(ROOT / 'starter-repo-blueprint' / '.env.example', r'''
NODE_ENV=development
APP_URL=http://localhost:3000
DATABASE_URL=postgresql://capsicum:capsicum@postgres:5432/capsicum
SESSION_SECRET=replace-with-at-least-32-random-bytes
OBJECT_STORAGE_ENDPOINT=http://minio:9000
OBJECT_STORAGE_REGION=us-east-1
OBJECT_STORAGE_BUCKET=capsicum
OBJECT_STORAGE_ACCESS_KEY=minio
OBJECT_STORAGE_SECRET_KEY=replace-me
PYTHON_WORKER_URL=http://worker-python:8000
R_WORKER_URL=http://worker-r:8001
AI_PROVIDER=mock
AI_MODEL=
AI_API_KEY=
EMBEDDING_PROVIDER=local
EMBEDDING_MODEL=
FEATURE_AI=true
FEATURE_VISION=true
FEATURE_ADVANCED_SIMULATIONS=true
FEATURE_QUANTITATIVE_MODELS=false
FEATURE_GENOMIC_MODELS=false
OTEL_EXPORTER_OTLP_ENDPOINT=
''')

write(ROOT / 'starter-repo-blueprint' / 'VERSION_POLICY.md', r'''
# Version Policy

- Use Node.js 24 LTS.
- Use Next.js 16.2 with the latest compatible stable patch.
- Use PostgreSQL 18 with the current stable minor release.
- Use AI SDK 6 with compatible provider packages.
- Pin exact dependency versions in the lockfile.
- Pin container images by version and preferably digest for release.
- Record Python and R package lock files.
- Do not use beta database or framework versions in the production profile.
''')

# Checklists
write(ROOT / 'checklists' / 'SCIENTIFIC_RELEASE_GATE.md', r'''
# Scientific Release Gate

- [ ] Claim is narrow and source-linked.
- [ ] Species, population, material, and method are identified.
- [ ] Allele/variant nomenclature is normalized and versioned.
- [ ] Applicability and exclusions are explicit.
- [ ] Conflicting evidence is recorded.
- [ ] Evidence grade is assigned independently from prediction eligibility.
- [ ] Rule has software fixtures and scientific review.
- [ ] Prohibited interpretations are displayed.
- [ ] Catalog release is immutable and reproducible.
- [ ] Historical runs remain linked to their original release.
''')

write(ROOT / 'checklists' / 'WEB_RELEASE_GATE.md', r'''
# Web Release Gate

- [ ] All required routes exist and are authorized.
- [ ] Core workflows pass E2E tests.
- [ ] Desktop, tablet, and mobile layouts pass visual review.
- [ ] Keyboard navigation and focus are complete.
- [ ] Automated accessibility scan has no blocking findings.
- [ ] Loading, empty, error, degraded, and denied states exist.
- [ ] Repeated clicks cannot create duplicate records or jobs.
- [ ] Scientific authority and uncertainty labels are visible.
- [ ] Unsupported predictions abstain.
- [ ] Export and backup/restore workflows are proven.
''')

write(ROOT / 'checklists' / 'AI_MODEL_RELEASE_GATE.md', r'''
# AI and Model Release Gate

- [ ] Training/evaluation dataset is versioned and licensed.
- [ ] Leakage audit is complete.
- [ ] Holdouts reflect family/accession, device, environment, and time as applicable.
- [ ] Model card and applicability are complete.
- [ ] Calibration and uncertainty are evaluated.
- [ ] Bias/error slices are reported.
- [ ] OOD or abstention behavior is tested.
- [ ] Security and prompt-injection tests pass.
- [ ] Human correction and rollback are available.
- [ ] Promotion approval is recorded.
''')

write(ROOT / 'DECISIONS_AND_STOP_CONDITIONS.md', r'''
# Decisions and Stop Conditions

## Immutable decisions

- The user interface is web-based and responsive.
- PostgreSQL is the authoritative relational store.
- The architecture is a modular monorepo with isolated Python/R workers, not a broad microservice estate.
- Exact genetics is implemented in a pure deterministic TypeScript package using rational probabilities.
- The supplied workbook is imported into normalized draft records; it is not the runtime schema or executable model.
- AI is an assistant and tool orchestrator, not a genetics authority.
- Vision outputs require protocol, quality, model, and human-review provenance.
- Quantitative/genomic capabilities are implemented as governed frameworks and remain unavailable until data/model gates pass.

## Stop and report instead of improvising when

- repository evidence or dependency compatibility prevents the approved architecture;
- a requested feature would require fabricating scientific accuracy;
- secure workspace isolation cannot be proven;
- an authoritative mutation cannot be audited;
- a schema migration risks irreversible data loss without a tested path;
- a worker contract cannot be validated across runtimes;
- required tests cannot meaningfully prove correctness;
- implementation would require weakening explicit nonclaims or approval gates.

## Do not stop merely because

- external AI credentials are absent;
- learned models or genomic datasets are absent;
- a scientific claim is still under review;
- object storage or email must use a local emulator;
- quantitative model accuracy cannot yet be validated.

In those cases, implement the platform, contracts, UI states, mocks, feature flags, and abstention behavior.
''')

write(ROOT / 'WORKSTREAM_CARDS.md', r'''
# Workstream Cards

## Workstream 0 — Integration lead and architecture

Mission: initialize the monorepo, enforce boundaries, own shared contracts, sequence merges, and run final integrated validation.

Allowed scope: root config, CI, shared contracts, dependency graph, integration fixes, final documentation.

Stop conditions: unresolved contract conflict, unsafe dependency cycle, or weakening of scientific/security gates.

## Workstream 1 — Platform, auth, data, and operations

Mission: identity, workspaces, roles, PostgreSQL, migrations, storage, jobs, audit, idempotency, observability, backups, and Docker Compose.

Owns: `auth`, `database`, `jobs`, `storage`, `observability`, config, infra.

Validation: migrations, authorization, isolation, concurrency, queue retry, backup/restore, compose smoke.

## Workstream 2 — Scientific catalog and import

Mission: normalize the supplied catalog, preserve provenance, implement evidence review, releases, applicability, and rule eligibility.

Owns: `scientific-catalog`, importer, catalog UI, source/review UI.

Must not: activate claims without the approved workflow.

Validation: zero silent-loss import, source resolution, release immutability, state transition tests.

## Workstream 3 — Genetics and simulation engines

Mission: implement exact genetics and advanced model contracts with mathematical proof and reproducibility.

Owns: `genetics-core`, `genetics-advanced`, simulation domain portions.

Must not: use AI or database calls in the exact engine.

Validation: golden, property-based, independent parity, performance thresholds.

## Workstream 4 — Breeding ledger and phenotyping records

Mission: germplasm, seed lots, plants, labels, crosses, harvests, progeny, pedigree, experiments, observations, selections, exports.

Owns: breeding and observation domain packages and corresponding web routes.

Validation: complete breeding-cycle E2E, pedigree constraints, QR/label identity, export round trip.

## Workstream 5 — Web experience and design system

Mission: accessible responsive interface, navigation, guided workflows, result visualization, operational states, and end-to-end UX integration.

Owns: `apps/web`, `ui` except shared backend contracts.

Validation: Playwright, accessibility, responsive visual review, repeated-click behavior, degraded states.

## Workstream 6 — AI research and copilot

Mission: document ingestion, retrieval, candidate extraction, tool orchestration, citations, evals, and user-facing copilot.

Owns: `ai`, research document pipelines, AI UI.

Must not: create authoritative rules or verified genotypes.

Validation: mock provider tests, citation evals, prompt injection, cross-workspace isolation, abstention.

## Workstream 7 — Vision and scientific workers

Mission: capture protocols, quality gates, deterministic measurements, annotation, model registry, Python/R workers, forward simulation templates.

Owns: `vision`, worker apps, worker contracts after integration-lead approval.

Validation: contract tests, measurement fixtures, job idempotency, model gating, worker failure recovery.

## Merge order

1. Workstream 0 repository baseline.
2. Workstream 1 platform contracts and schema foundations.
3. Workstream 2 normalized catalog plus Workstream 3 pure engine in parallel.
4. Workstream 4 breeding domain.
5. Workstream 5 integrated web workflows.
6. Workstream 6 AI on stable catalog/simulation APIs.
7. Workstream 7 vision/workers on stable material/observation/job contracts.
8. Workstream 0 integrated hardening and final validation.

Only the integration lead edits root lockfiles, shared generated contracts, global migrations sequencing, and root CI after initialization.
''')

write(ROOT / 'VALIDATION_LEDGER.md', r'''
# Validation Ledger

The coding agent must replace command placeholders with discovered repository commands and record exact outcomes.

| Layer | Required proof | Blocking |
|---|---|---:|
| Repository | clean install from lockfile; task graph works | Yes |
| Formatting/lint | no introduced errors or material warnings | Yes |
| Type safety | strict TypeScript, Python type/lint, R checks | Yes |
| Contracts | TypeScript/Python/R contract parity | Yes |
| Database | empty install, upgrade migration, constraints, seed import | Yes |
| Auth/security | workspace isolation, permissions, CSRF/session, file safety | Yes |
| Exact genetics | golden, property, independent parity, performance | Yes |
| Advanced genetics | linkage/maternal/uncertainty fixtures and invariants | Yes for released engines |
| Web | component/integration/E2E, responsive, accessibility | Yes |
| Catalog | reconciliation, review transitions, immutable release | Yes |
| Breeding ledger | complete breeding-cycle E2E and pedigree integrity | Yes |
| AI | mock provider, citations, tool policy, injection, isolation | Yes if feature enabled |
| Vision | capture quality, deterministic measurement, correction flow | Yes if feature enabled |
| Workers/jobs | idempotency, retry, timeout, cancellation, artifacts | Yes |
| Build/deploy | production builds and Compose smoke | Yes |
| Operations | backup/restore, health, logs, metrics, export | Yes |
| Scientific claims | model-specific release gate | Yes for each promoted claim |
''')

write(ROOT / 'FINAL_REPORT_TEMPLATE.md', r'''
# Mandatory Final Agent Report

## 1. Executive result

- Completion status: complete / conditionally complete / blocked
- Repository path and commit:
- Deployable services:
- Primary remaining scientific gates:

## 2. Pre-edit/current-state findings

- Source pack conflicts:
- Decisions resolved:
- Dependency versions selected and why:
- Deviations from the proposed topology:

## 3. Implementation by workstream

| Workstream | Packages/routes/services created | Outcome |
|---|---|---|

## 4. Database and data migration

- Schema/migrations:
- Catalog import reconciliation totals:
- Seed/demo data:
- Data-loss or compatibility risks:

## 5. Scientific and model integrity

- Executable rules:
- Unsupported/disabled models:
- Exact engine parity evidence:
- AI/vision/model authority boundaries:

## 6. Web application

- Route inventory:
- Core E2E workflows:
- Accessibility and responsive evidence:
- Error/degraded states:

## 7. Validation evidence

| Command/check | Scope | Result | Evidence/log location |
|---|---|---|---|

Classify every failure as pre-existing, introduced, environment/tooling, or unresolved.

## 8. Security and operations

- Auth/workspace isolation:
- Secrets/file handling:
- Dependency/container scan:
- Backups and restore result:
- Observability:
- Rollout and rollback:

## 9. Acceptance criteria traceability

Map every criterion in `docs/12_DEFINITION_OF_DONE.md` to pass/fail/blocked and concrete evidence.

## 10. Residual risks and debt

For each item include severity, impact, owner area, removal path, and whether it blocks release.

## 11. Files changed

Provide a complete grouped file list.
''')

# Master execution prompt
write(ROOT / 'EXECUTION_PROMPT.md', r'''
# One-Shot Autonomous Coding-Agent Execution Prompt

You are the principal engineer and integration owner responsible for creating the complete **Capsicum Breeding Intelligence Platform** from this handoff pack. This is a greenfield, multi-language, web-based scientific application. Execute the work end-to-end. Use bounded sub-agents or workstreams where helpful, but retain one integration owner and do not finish until integrated validation is complete or a genuine stop condition is reached.

## Mission

Create a production-quality monorepo and runnable web application that enables Capsicum breeders to manage biological material and pedigrees, record controlled crosses and observations, run scientifically governed genetic simulations, use AI for evidence-backed research and workflow assistance, and perform controlled visual phenotype analysis.

The application must be useful without AI credentials, learned vision models, genomic datasets, or unreviewed scientific claims. Unsupported capabilities must abstain or remain feature-gated rather than returning fabricated predictions.

## Source pack to obey

Read the entire extracted pack before editing. The primary source is:

- `source-pack/CAPSICUM_BREEDING_INTELLIGENCE_PLATFORM_IMPLEMENTATION_PLAN.md`

Also obey:

- `DECISIONS_AND_STOP_CONDITIONS.md`
- `docs/*.md`
- `contracts/*.schema.json`
- `WORKSTREAM_CARDS.md`
- `VALIDATION_LEDGER.md`
- the supplied workbook, CSVs, and source registry under `source-pack/` and `data/`

Readiness is **conditionally ready**: the platform and exact/advanced simulation architecture can be implemented. Data-trained quantitative, genomic, and learned vision claims cannot be promoted without adequate data and validation.

## Execution mode

Primary mode: greenfield platform build.

Secondary obligations: scientific governance, security hardening, data platform, AI safety, computer vision, multi-agent coordination, accessibility, and local-first operations.

## Working behavior

- Inspect the entire source pack before creating the repository.
- Distinguish facts, requirements, implementation decisions, assumptions, and scientific blockers.
- Do not ask routine questions. Make defensible implementation decisions consistent with the pack and document them.
- Do not weaken scope because the project is large. Implement in release slices and integrate them.
- Do not claim scientific accuracy where no validated data/model exists.
- Do not stop merely because hosted credentials or training datasets are absent. Use local emulators, mocks, feature flags, and `not_validated` states.
- Use sub-agents with nonoverlapping scopes from `WORKSTREAM_CARDS.md`.
- One integration owner controls root configuration, lockfile, shared contracts, migration order, and final validation.

## Immutable architecture

Create a modular monorepo containing:

- Next.js 16.2 App Router web application on Node.js 24 LTS;
- PostgreSQL 18 current stable minor;
- S3-compatible object storage, with MinIO locally;
- PostgreSQL-backed jobs with idempotency, retries, timeouts, cancellation, and observability;
- pure TypeScript exact genetics package;
- advanced TypeScript genetics package;
- normalized scientific catalog and approval workflows;
- breeding, pedigree, experiment, phenotype, simulation, model, research, and audit domains;
- provider-agnostic AI SDK 6 integration with deterministic mock provider tests;
- Python worker for image/scientific processing;
- R worker for AlphaSimR and approved statistical jobs;
- Docker Compose local deployment;
- CI, backups, restore testing, documentation, and demo data.

Do not create a broad microservice architecture. Do not place deterministic genetics inside an LLM, Python notebook, React component, or database trigger.

## Package boundary rules

- `genetics-core` is pure, deterministic, exact, and has no network, database, React, worker, or AI dependency.
- `genetics-advanced` depends on core primitives and approved model contracts.
- `scientific-catalog` owns evidence and executable-rule eligibility.
- AI depends on catalog/search/simulation tools; no authoritative package depends on AI.
- Web owns authorization and user orchestration.
- Workers consume versioned contracts and cannot trust user-supplied workspace or permission fields.
- Database access is scoped through domain repositories/services. No generic unscoped CRUD.
- Generated types and clients are changed through their source contract.

## Scientific invariants

- Never infer dominance from capitalization.
- Never silently substitute unknown genotype input.
- Preserve maternal and paternal direction.
- Separate genotype segregation from phenotype interpretation.
- Separate conception from surviving-offspring probabilities when lethal rules apply.
- Store assumptions, evidence state, applicability, exclusions, and versions with every result.
- A cultivar name is not a genotype.
- AI cannot verify a genotype or approve a claim.
- Quantitative and genomic predictions require promoted models and compatible data.
- Disease resistance requires pathogen context.
- Image-derived observations require capture, quality, method/model, and correction provenance.

## Required pre-implementation inspection output

Before edits, produce a concise internal execution memo recording:

- files in this pack inspected;
- final repository topology;
- selected stable dependency versions;
- ORM/query/migration approach and how PostgreSQL-native constraints remain available;
- auth architecture;
- job/worker transport;
- object storage design;
- shared contract generation/validation strategy;
- package dependency direction;
- release slice and sub-agent allocation;
- any conflict with the source pack.

Do not wait for approval unless a genuine stop condition is present.

## Implementation sequence

### 1. Repository and platform foundation

- Initialize pnpm workspace and task graph.
- Create web, Python worker, and R worker deployables.
- Establish strict TypeScript, lint, formatting, tests, CI, and container builds.
- Implement typed environment validation.
- Implement identity, workspaces, memberships, roles, permissions, sessions, and audit.
- Implement PostgreSQL schema/migrations and object storage.
- Implement PostgreSQL-backed jobs, transactional outbox, idempotency, retries, cancellation, and admin job UI.
- Implement Docker Compose, health checks, backup scripts, and restore test.

### 2. Scientific catalog foundation

- Import the supplied workbook through staging and reconciliation.
- Preserve raw source payload, source file hash, source sheet/row, and source links.
- Normalize loci, aliases, evidence assertions, publications, applicability, review state, prohibited claims, and releases.
- Implement review and approval state machine with separation of duties.
- Seed a draft catalog release and demo workspace.
- Do not convert spreadsheet descriptions directly into executable rules.

### 3. Exact and advanced simulations

- Implement rational arithmetic and unique weighted gametes.
- Implement exact nuclear, multi-allelic, independent multi-locus inheritance.
- Implement versioned phenotype rules, penetrance, and lethal distributions.
- Implement unknown/probabilistic genotype marginalization.
- Implement target recovery and operational adjustment.
- Implement linkage/phase, maternal/cytoplasmic inheritance, epistasis rule graphs, host–pathogen contexts, and Monte Carlo fallback.
- Store immutable runs and results using the supplied contracts.
- Build golden fixtures, property tests, independent parity engine/tests, and performance guards.

### 4. Breeding ledger

- Implement taxa, germplasm, seed lots, plants, material events, labels/QR, crosses, pollination, verification, fruit, seed harvest, progeny, generation, pedigree, selection goals, decisions, experiments, variables, sessions, observations, environments, and exports.
- Reject pedigree cycles and conflicting origin events.
- Build complete guided web workflows and a breeding-cycle E2E test.

### 5. Web application

- Implement the route map in `docs/02_WEB_APPLICATION_UX_SPEC.md`.
- Build a professional responsive design system and accessible application shell.
- Implement dashboard, table/detail/create/edit flows, guided cross planner, simulation results, pedigree visualization, phenotype capture, catalog review, AI, reports, settings, and admin operations.
- Include all loading, empty, partial, error, degraded, permission, and unsupported states.
- Prevent duplicate actions on repeated clicks at both UI and server levels.
- Use server components by default and client components only where interaction requires them.

### 6. AI research and copilot

- Implement versioned research document ingestion and source passages.
- Implement hybrid retrieval over approved/draft evidence with explicit state filters.
- Implement structured candidate extraction and review queue.
- Implement bounded AI tools for evidence search, simulation, record access, and draft creation.
- Implement cited AI responses using `contracts/ai-answer.schema.json`.
- Use a deterministic mock provider for CI and permit hosted providers through environment configuration.
- Implement evaluation, prompt-injection defense, workspace isolation, and provider degradation.

### 7. Phenotype vision and workers

- Implement capture protocols, image upload, object storage, quality gates, calibration, deterministic measurements, annotation revisions, and human correction.
- Implement a Python worker contract and a baseline deterministic fruit-image pipeline with documented fixtures.
- Implement a model registry, dataset registry, evaluations, promotion gates, and `not_validated` states.
- Support optional annotation assistance but do not promote unvalidated zero-shot output.
- Implement an R worker and reproducible AlphaSimR example jobs; keep outputs research-only and distinct from exact simulation authority.

### 8. Statistical/model framework

- Implement training dataset, model version, model card, evaluation, promotion, rollback, and feature-availability workflows.
- Provide interfaces and reproducible example/synthetic tests for mixed, G×E, genomic, and selection models.
- Keep user-facing production predictions disabled until explicit data and validation gates pass.

### 9. Interoperability, reports, and operations

- Implement MIAPPE 1.2 export, MCPD-compatible germplasm export, and selected BrAPI 2.1 endpoints after internal contracts stabilize.
- Implement CSV/JSON exports and full workspace scientific-data export.
- Complete observability, admin health, backup/restore, security headers, upload safety, dependency scanning, and runbooks.

### 10. Integrated hardening

- Run the entire validation ladder.
- Fix all introduced failures and material warnings.
- Verify clean bootstrap from a fresh clone and empty database.
- Verify catalog import reconciliation.
- Verify Docker Compose and production builds.
- Verify core web E2E workflows, accessibility, responsive layouts, repeated-click behavior, AI degraded state, worker retry, and backup restore.
- Complete the mandatory final report.

## Required acceptance outcomes

At minimum, the final repository must prove:

1. Complete web interface and route inventory.
2. Onboarding and workspace roles.
3. Catalog import with reconciliation and governed release.
4. Exact supported simulation with evidence/assumptions/versions.
5. Advanced engines and abstention behavior.
6. Full breeding cycle and pedigree.
7. Standardized observations and phenotype image workflow.
8. AI research/citation/tool flow with mock provider.
9. Worker jobs and research-only simulation/model infrastructure.
10. Exports, audit, observability, backup, restore, Docker Compose, CI, security, accessibility, and E2E proof.

## Validation sequence

Discover and record exact commands. The completed project must include and pass equivalents of:

- dependency install with frozen lockfile;
- format check;
- lint/static analysis;
- strict TypeScript typecheck;
- Python lint/type/test;
- R check/test;
- JSON/OpenAPI/contract validation;
- database migration and seed/import tests;
- package unit tests;
- genetics golden/property/parity tests;
- domain integration tests;
- web component and server integration tests;
- Playwright E2E tests;
- accessibility scans;
- production builds;
- security/dependency/secret scans;
- Docker Compose smoke;
- backup and restore drill;
- export/import round trip.

Classify failures honestly. Do not suppress warnings that affect correctness, security, science, accessibility, performance, or release quality.

## Stop conditions

Stop and report only if:

- the source pack contains an irreconcilable contradiction affecting architecture, security, data, or scientific authority;
- compatible stable dependencies cannot implement the required architecture;
- secure workspace isolation or data integrity cannot be proven;
- required implementation would fabricate a scientific claim;
- validation cannot meaningfully prove correctness;
- completion requires unauthorized external data or credentials and no mock/feature-gated implementation is possible.

Do not stop for ordinary implementation complexity, absent provider keys, absent training data, or scientific claims that should remain disabled.

## Required final response

Use `FINAL_REPORT_TEMPLATE.md`. Include exact commands, outcomes, complete file inventory, catalog reconciliation counts, routes, deployables, feature flags, model/claim states, residual risks, and acceptance-criteria traceability.

## Completion standard

Do not end after scaffolding, partial CRUD, or isolated packages. Complete the integrated web application and supporting services to the maximum scientifically honest extent possible. A capability may be marked research-only or unavailable due to data, but its governance, contracts, UI state, and future integration boundary must be implemented. The result must start locally, build for production, pass the required validation, and preserve every scientific and operational boundary in this pack.
''')

# Additional pack inventory/documentation.
write(ROOT / 'PACKAGE_CONTENTS.md', r'''
# Package Contents

## Agent control files

- `START_HERE.md`
- `EXECUTION_PROMPT.md`
- `SOURCE_PACK_READINESS.md`
- `DECISIONS_AND_STOP_CONDITIONS.md`
- `WORKSTREAM_CARDS.md`
- `VALIDATION_LEDGER.md`
- `FINAL_REPORT_TEMPLATE.md`

## Specifications

- Product charter
- System architecture
- Web application UX
- Data model and migrations
- Simulation engines
- AI/RAG
- Phenotype vision
- Scientific governance
- APIs/contracts
- Security and operations
- Testing
- Delivery/release
- Definition of done
- Environment/deployment
- Repository conventions
- Verified technology/standards baseline

## Evidence and data

- Original workbook and source CSVs
- Original catalog generator
- Full implementation plan
- Every workbook sheet exported as CSV and JSON
- Workbook export manifest

## Machine-readable contracts

- simulation request/result
- AI answer
- phenotype observation
- catalog import record

## Blueprints and checklists

- expected repository tree
- root agent rules
- environment template
- version policy
- scientific, web, and model release gates
''')

# Include generator for reproducibility.
shutil.copy2('/mnt/data/build_capsicum_handoff_pack.py', ROOT / 'tools' / 'build_capsicum_handoff_pack.py')

# Build manifest and checksums, excluding manifest files themselves on first pass.
manifest_entries = []
for path in sorted(ROOT.rglob('*')):
    if not path.is_file() or path.name in {'MANIFEST.json', 'SHA256SUMS.txt'}:
        continue
    data = path.read_bytes()
    manifest_entries.append({
        'path': path.relative_to(ROOT).as_posix(),
        'sizeBytes': len(data),
        'sha256': hashlib.sha256(data).hexdigest(),
    })
manifest = {
    'package': 'capsicum-breeding-intelligence-platform-agent-handoff',
    'version': '1.0.0',
    'createdDate': '2026-07-22',
    'readiness': 'conditionally-ready',
    'fileCount': len(manifest_entries),
    'files': manifest_entries,
}
dump_json(ROOT / 'MANIFEST.json', manifest)
write(ROOT / 'SHA256SUMS.txt', '\n'.join(f"{item['sha256']}  {item['path']}" for item in manifest_entries))

# Zip with root folder included.
zip_path = Path('/mnt/data/capsicum-breeding-intelligence-platform-agent-handoff-v1.zip')
if zip_path.exists():
    zip_path.unlink()
with zipfile.ZipFile(zip_path, 'w', compression=zipfile.ZIP_DEFLATED, compresslevel=9) as zf:
    for path in sorted(ROOT.rglob('*')):
        if path.is_file():
            arc = Path(ROOT.name) / path.relative_to(ROOT)
            zf.write(path, arcname=arc.as_posix())

print(f'Created {zip_path}')
print(f'Files: {sum(1 for p in ROOT.rglob("*") if p.is_file())}')
print(f'Zip bytes: {zip_path.stat().st_size}')
